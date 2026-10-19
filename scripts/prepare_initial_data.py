"""Prepare a private migration backup. This script never uploads business data."""
import argparse
import json
import os
import sys
import tempfile
import re
import io
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent.parent))
from backend.database import initialize,connection,now
from backend.transfers import preview_import,commit_import
from backend.main import BACKUP_TABLES
from openpyxl import load_workbook


def business_from_workbook(source):
    """Read the supplied workbook's business block; never hardcode private details."""
    workbook=load_workbook(io.BytesIO(source),read_only=True,data_only=True,keep_links=False)
    sheet=next((s for s in workbook if s.title.lower()=='invoice'),None)
    if sheet is None:
        workbook.close();return {}
    rows=list(sheet.iter_rows(max_row=6,max_col=7,values_only=True))
    def cell(row,col=0):
        v=rows[row-1][col]
        return str(int(v)) if isinstance(v,(int,float)) and int(v)==v else str(v or '').strip()
    address='\n'.join(filter(None,[cell(3),cell(4)]))
    pin=re.search(r'\b\d{6}\b',address)
    gstin=re.sub(r'^GSTIN\s*:?\s*','',cell(5),flags=re.I)
    phone=re.sub(r'^Phone\s*[:\-]*\s*','',cell(6),flags=re.I)
    result={'business_name':cell(2),'address':address,'gstin':gstin,'phone':' / '.join(filter(None,[phone,cell(6,1)])),'pincode':pin.group(0) if pin else ''}
    last=address.split(',')[-1].strip()
    result['city']=re.sub(r'[\s-]*\d{6}.*$','',last).strip()
    workbook.close()
    return {k:v for k,v in result.items() if v}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook',type=Path)
    parser.add_argument('--output',required=True,type=Path)
    parser.add_argument('--review-output',required=True,type=Path)
    parser.add_argument('--allow-duplicate-customers',action='store_true')
    parser.add_argument('--historical-unpaid',action='store_true')
    args=parser.parse_args()
    source=args.workbook.read_bytes()
    review={"source":args.workbook.name,"created_at":now(),"notes":["Product rates are not inferred from descriptions.","Historical invoices lack reliable tax and payment details.","Unresolved duplicate customer matches are not linked automatically."],"imports":{}}
    with tempfile.TemporaryDirectory(prefix='svs-migration-') as tmp:
        os.environ['DATA_DIR']=tmp;initialize()
        from backend.service import update_settings
        update_settings(business_from_workbook(source),'Owner-approved workbook migration')
        for kind in ['customers','items','orders']:
            preview=preview_import(source,args.workbook.name,kind)
            selected=[r['index'] for r in preview['records'] if not r['error'] and (kind!='orders' or args.historical_unpaid) and (kind!='customers' or args.allow_duplicate_customers or not any('duplicate' in w.lower() for w in r['warnings']))]
            result=commit_import(preview['id'],selected,'Owner-approved workbook migration') if selected else {'imported':0}
            review['imports'][kind]={"imported":result['imported'],"valid":preview['valid'],"incomplete":preview['incomplete'],"blank_rows_scanned":preview['skipped'],"tail_not_scanned":preview['truncated'],"review_rows":[r for r in preview['records'] if r['error'] or (kind=='customers' and r['warnings'])]}
        with connection() as db:
            backup={"format":"svs-billing-backup","version":1,"created_at":now(),"currency":"INR","money_unit":"paise","tables":{t:[dict(r) for r in db.execute('SELECT * FROM '+t)] for t in BACKUP_TABLES}}
        for path,value in [(args.output,backup),(args.review_output,review)]:
            path.parent.mkdir(parents=True,exist_ok=True)
            path.write_text(json.dumps(value,ensure_ascii=False,indent=2));path.chmod(0o600)
        print(json.dumps({kind:data['imported'] for kind,data in review['imports'].items()}))


if __name__=='__main__':
    main()
