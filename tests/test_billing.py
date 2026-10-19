import io
import json
import uuid
from concurrent.futures import ThreadPoolExecutor
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from backend.main import app
from backend.database import initialize, connection
from backend.auth import set_owner
from backend.money import calculate, amount_words


@pytest.fixture(params=['sqlite','turso-protocol'])
def client(tmp_path,monkeypatch,request):
    for key in ('DB_BACKEND','TURSO_DATABASE_URL','TURSO_AUTH_TOKEN'):
        monkeypatch.delenv(key,raising=False)
    if request.param=='turso-protocol':
        request.getfixturevalue('turso_backend')
    monkeypatch.setenv('DATA_DIR',str(tmp_path))
    monkeypatch.setenv('APP_ENV','development')
    monkeypatch.delenv('PUBLIC_URL',raising=False)
    monkeypatch.delenv('RENDER_EXTERNAL_URL',raising=False)
    monkeypatch.delenv('OWNER_PASSWORD',raising=False)
    initialize();set_owner('owner@example.test','Test password 2026!')
    with TestClient(app) as c:
        response=c.post('/api/auth/login',json={'email':'owner@example.test','password':'Test password 2026!'})
        assert response.status_code==200
        c.headers['X-CSRF-Token']=response.json()['csrf']
        yield c


def create_customer(c,name='Test Beauty Store'):
    r=c.post('/api/customers',json={'name':name,'place':'Attibele','gstin':'29TEST0000A1Z1'})
    assert r.status_code==200,r.text
    return r.json()


def invoice_payload(c,**overrides):
    customer=create_customer(c)
    return {'customer_id':customer['id'],'invoice_date':'2026-10-19','lines':[{'description':'Nisha test product','quantity':'80','rate':'77','gst_rate':'18'}],'finalize':True,**overrides}


def test_excel_acceptance_and_rounding():
    data=calculate([{'description':str(i),'quantity':str(q),'rate':str(rate)} for i,(q,rate) in enumerate([(50,77),(80,77),(20,77),(14,190),(192,'12.50'),(25,49)])])
    assert [x['taxable'] for x in data['lines']]==[326271,522034,130508,225424,203390,103814]
    assert data['taxable']==1511441
    assert data['cgst']==136030 and data['sgst']==136030
    assert data['rounding']==-1 and data['total']==1783500
    assert data['taxable']+data['cgst']+data['sgst']+data['rounding']==data['total']
    assert amount_words(data['total'])=='Rupees Seventeen Thousand Eight Hundred Thirty Five Only'
    assert amount_words(123456789)=='Rupees Twelve Lakh Thirty Four Thousand Five Hundred Sixty Seven and Eighty Nine Paise Only'


@pytest.mark.parametrize('value',['NaN','Infinity','-1','not a price'])
def test_invalid_financial_input(client,value):
    r=client.post('/api/calculate',json={'lines':[{'description':'Product','quantity':'1','rate':value}]})
    assert r.status_code==400


def test_auth_csrf_and_origin(client):
    token=client.headers.pop('X-CSRF-Token')
    assert client.post('/api/customers',json={'name':'Unauthorized change'}).status_code==403
    client.headers['X-CSRF-Token']=token
    assert client.post('/api/customers',headers={'Origin':'https://hostile.example'},json={'name':'Cross-site change'}).status_code==403
    assert client.post('/api/auth/login',headers={'Origin':'https://hostile.example'},json={'email':'owner@example.test','password':'Test password 2026!'}).status_code==403
    assert client.get('/api/settings').headers['cache-control']=='no-store'
    assert "script-src 'self'" in client.get('/').headers['content-security-policy']
    assert client.post('/api/auth/logout',json={}).status_code==200
    assert client.get('/api/invoices').status_code==401


def test_persistent_login_throttle(client):
    for _ in range(10):
        assert client.post('/api/auth/login',json={'email':'unknown@example.test','password':'Wrong password!'}).status_code==401
    assert client.post('/api/auth/login',json={'email':'owner@example.test','password':'Test password 2026!'}).status_code==429


def test_draft_finalization_and_snapshots(client):
    payload=invoice_payload(client,finalize=False)
    draft=client.post('/api/invoices',json=payload).json()
    assert draft['number'] is None and draft['status']=='Draft'
    result=client.put('/api/invoices/'+draft['id'],json={**payload,'finalize':True,'version':draft['version']})
    assert result.status_code==200,result.text
    invoice=result.json();assert invoice['number']=='1'
    client.put('/api/customers/'+payload['customer_id'],json={'name':'Changed customer name'})
    client.put('/api/settings',json={'business_name':'Changed company'})
    old=client.get('/api/invoices/'+invoice['id']).json()
    assert old['customer_snapshot']['name']=='Test Beauty Store'
    assert old['business_snapshot']['business_name']=='SVS ENTERPRISES'
    changed=client.put('/api/invoices/'+old['id'],json={**payload,'finalize':True,'version':old['version'],'notes':'Correction'})
    assert changed.status_code==200,changed.text
    assert changed.json()['customer_snapshot']['name']=='Test Beauty Store'
    assert changed.json()['business_snapshot']['business_name']=='SVS ENTERPRISES'
    assert client.put('/api/invoices/'+old['id'],json={**payload,'version':1}).status_code==400


def test_unique_numbering_under_concurrent_saves(client):
    payload=invoice_payload(client)
    def save(_):
        response=client.post('/api/invoices',json=payload)
        assert response.status_code==200,response.text
        return response.json()['number']
    with ThreadPoolExecutor(max_workers=5) as pool:
        numbers=list(pool.map(save,range(15)))
    assert len(set(numbers))==15
    assert sorted(map(int,numbers))==list(range(1,16))


def test_financial_year_reset_stays_unique(client):
    client.put('/api/settings',json={'invoice_prefix':'SVS/','reset_financial_year':True,'number_padding':3})
    payload=invoice_payload(client)
    first=client.post('/api/invoices',json={**payload,'invoice_date':'2026-03-31'}).json()
    second=client.post('/api/invoices',json={**payload,'invoice_date':'2026-04-01'}).json()
    assert first['number']=='SVS/2025-26/001'
    assert second['number']=='SVS/2026-27/001'


def test_payments_overpayment_edit_void_and_cancel(client):
    invoice=client.post('/api/invoices',json=invoice_payload(client)).json()
    key=str(uuid.uuid4())
    data={'invoice_id':invoice['id'],'payment_date':'2026-10-19','amount':'1000','method':'UPI','idempotency_key':key}
    first=client.post('/api/payments',json=data);assert first.status_code==200,first.text
    assert first.json()['status']=='Partially Paid'
    assert first.json()['balance']==516000
    again=client.post('/api/payments',json=data).json()
    assert len(again['payments'])==1
    assert client.post('/api/payments',json={**data,'amount':'6000','idempotency_key':str(uuid.uuid4())}).status_code==400
    assert client.post('/api/invoices/'+invoice['id']+'/cancel',json={'reason':'Incorrect invoice'}).status_code==400
    edit=client.put('/api/payments/'+key,json={**data,'amount':'6160'})
    assert edit.status_code==200 and edit.json()['status']=='Paid'
    void=client.post('/api/payments/'+key+'/void',json={'reason':'Payment entered against wrong bill'})
    assert void.status_code==200 and void.json()['status']=='Unpaid'
    cancel=client.post('/api/invoices/'+invoice['id']+'/cancel',json={'reason':'Replaced with corrected invoice'})
    assert cancel.json()['status']=='Cancelled'
    assert client.get('/api/reports/sales?start=2026-01-01&end=2026-12-31').json()['summary']['total']==0
    actions=[r['action'] for r in client.get('/api/audit').json()['rows']]
    assert all(x in actions for x in ['Payment added','Payment edited','Payment voided','Invoice cancelled'])


def test_discount_is_distinct_from_legacy_label(client):
    line={'description':'Product','quantity':'2','rate':'118','discount_percent':'10'}
    assert client.post('/api/calculate',json={'lines':[line]}).status_code==400
    assert client.put('/api/settings',json={'enable_discounts':True,'legacy_labels':True}).status_code==200
    data=client.post('/api/calculate',json={'lines':[line]}).json()
    assert data['total']==21240 and data['discount']==2360 and data['taxable']==18000


def xlsx(sheets):
    wb=Workbook();wb.remove(wb.active)
    for name,rows in sheets.items():
        ws=wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    output=io.BytesIO();wb.save(output);return output.getvalue()


def preview(client,kind,data):
    return client.post('/api/import/preview',data={'kind':kind},files={'file':('source.xlsx',data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})


def test_import_review_duplicate_names_and_preserved_descriptions(client):
    data=xlsx({'Customer':[['Name','Place'],['Same shop','First street'],['Same shop','Second street']],
               'Item':[[None,'Nisha Black 15 Rupees'],[None,'Item Name'],[None,'Neeta cone']]})
    customers=preview(client,'customers',data).json()
    assert customers['valid']==2 and not customers['records'][1]['warnings']
    assert all(r['recommended'] for r in customers['records'])
    assert client.get('/api/customers').json()['total']==0
    commit=client.post('/api/import/commit',json={'preview_id':customers['id'],'selected':[0,1]})
    assert commit.status_code==200,commit.text
    assert client.get('/api/customers').json()['total']==2
    assert all('duplicate' not in r for r in client.get('/api/customers').json()['rows'])
    assert client.post('/api/import/commit',json={'preview_id':customers['id'],'selected':[0]}).status_code==400
    items=preview(client,'items',data).json()
    assert items['valid']==2 and items['incomplete']==1
    assert items['records'][0]['data']['default_rate'] is None
    client.post('/api/import/commit',json={'preview_id':items['id'],'selected':[0,2]})
    rows=client.get('/api/items').json()['rows']
    assert any(r['name']=='Nisha Black 15 Rupees' and r['default_rate'] is None for r in rows)


def test_historical_orders_skip_formulas_and_keep_tax_unknown(client):
    data=xlsx({'Orders':[['Sl.No','Customer Name','Details','Date','Amount'],[136,'Old customer','Old address','2019-01-14',2669.94],[137,None,'#N/A',None,None],[138,'Incomplete','','',None],[139,'Bad total','','2019-01-15','#N/A']]})
    result=preview(client,'orders',data);assert result.status_code==200,result.text
    p=result.json();assert p['valid']==1 and p['incomplete']==2 and p['skipped']==1
    result=client.post('/api/import/commit',json={'preview_id':p['id'],'selected':[0]})
    assert result.status_code==200,result.text
    row=client.get('/api/invoices').json()['rows'][0]
    d=client.get('/api/invoices/'+row['id']).json()
    assert d['historical']==1 and d['total']==266994 and d['taxable']==0 and d['tax_groups']==[]
    assert d['status']=='Unpaid' and d['balance']==266994
    report=client.get('/api/reports/sales?start=2019-01-01&end=2019-12-31').json()
    assert report['summary']['total']==266994 and report['summary']['cgst']==0


def test_pdf_multipage_and_safe_exports(client):
    payload=invoice_payload(client,lines=[{'description':'Long product description '+str(i),'quantity':'1','rate':'118'} for i in range(90)])
    invoice=client.post('/api/invoices',json=payload).json()
    pdf=client.get('/api/invoices/'+invoice['id']+'/pdf')
    assert pdf.status_code==200 and pdf.content.startswith(b'%PDF')
    reader=PdfReader(io.BytesIO(pdf.content));assert len(reader.pages)>1
    texts=[p.extract_text() for p in reader.pages]
    assert all('Description of goods' in t for t in texts[:-1])
    assert 'GRAND TOTAL' in texts[-1]
    create_customer(client,'=HYPERLINK("https://example.com")')
    export=client.get('/api/export/customers')
    sheet=load_workbook(io.BytesIO(export.content),data_only=False).active
    cells=[r[0] for r in sheet.iter_rows(min_row=2)]
    malicious=next(c for c in cells if 'HYPERLINK' in str(c.value))
    assert malicious.data_type!='f'
    report=client.get('/api/reports/sales?start=2026-01-01&end=2026-12-31&format=pdf')
    assert report.status_code==200 and report.content.startswith(b'%PDF')


def test_backup_relations_and_credentials_excluded(client):
    invoice=client.post('/api/invoices',json=invoice_payload(client)).json()
    backup=client.get('/api/backup').json()
    assert backup['format']=='svs-billing-backup'
    assert 'users' not in backup['tables'] and 'sessions' not in backup['tables']
    assert len(backup['tables']['invoices'])==1 and len(backup['tables']['invoice_items'])==1
    assert backup['tables']['invoice_items'][0]['invoice_id']==invoice['id']
    assert 'password_hash' not in json.dumps(backup)


def test_pagination_and_customer_history(client):
    payload=invoice_payload(client)
    for _ in range(3):
        assert client.post('/api/invoices',json=payload).status_code==200
    result=client.get('/api/invoices?page=2&page_size=2').json()
    assert result['total']==3 and len(result['rows'])==1
    profile=client.get('/api/customers/'+payload['customer_id']).json()
    assert profile['summary']['invoices']==3 and profile['summary']['outstanding']==1848000


def test_private_initial_data_import_is_reviewed_and_empty_only(client):
    source=client.get('/api/backup').json()
    source['tables']['customers']=[{'id':str(uuid.uuid4()),'name':'Migrated customer','place':'','city':'','state':'','pincode':'','phone':'','gstin':'','email':'','notes':'','active':1,'created_at':'2026-10-19T00:00:00+00:00','updated_at':'2026-10-19T00:00:00+00:00'}]
    response=client.post('/api/backup/preview',files={'file':('initial.json',json.dumps(source),'application/json')})
    assert response.status_code==200,response.text
    preview=response.json();assert preview['counts']['customers']==1
    assert client.get('/api/customers').json()['total']==0
    assert client.post('/api/backup/initialize',json={'preview_id':preview['id'],'confirmation':'yes'}).status_code==400
    result=client.post('/api/backup/initialize',json={'preview_id':preview['id'],'confirmation':'IMPORT'})
    assert result.status_code==200,result.text
    assert client.get('/api/customers').json()['total']==1
    assert client.get('/api/auth/me').json()['email']=='owner@example.test'
    assert client.post('/api/backup/preview',files={'file':('initial.json',json.dumps(source),'application/json')}).status_code==400


def test_ambiguous_historical_customers_do_not_collapse_in_reports(client):
    source=xlsx({'Orders':[['Sl.No','Customer Name','Date','Amount'],[1,'Unknown customer A','2019-01-01',118],[2,'Unknown customer B','2019-01-02',236]]})
    p=preview(client,'orders',source).json()
    result=client.post('/api/import/commit',json={'preview_id':p['id'],'selected':[0,1]})
    assert result.status_code==200,result.text
    data=client.get('/api/reports/customers?start=2019-01-01&end=2019-12-31').json()
    assert len(data['rows'])==2
    assert data['summary']['historical_invoices']==2
    export=client.get('/api/reports/sales?start=2019-01-01&end=2019-12-31&format=xlsx')
    rows=list(load_workbook(io.BytesIO(export.content),data_only=True).active.values)
    taxable_index=rows[0].index('Taxable')
    assert rows[1][taxable_index] is None


def test_delete_masters_keeps_invoice_snapshots_payments_and_editing(client):
    payload=invoice_payload(client,invoice_date='2026-10-07')
    item=client.post('/api/items',json={'name':'Saved product','default_rate':'118'}).json()
    payload['lines']=[{'item_id':item['id'],'description':'Original product description','quantity':'2','rate':'118','gst_rate':'18'}]
    invoice=client.post('/api/invoices',json=payload).json()
    payment={'invoice_id':invoice['id'],'payment_date':'2026-10-08','amount':'100','method':'Cash','idempotency_key':str(uuid.uuid4())}
    before=client.post('/api/payments',json=payment).json()
    for kind,identifier in [('customers',payload['customer_id']),('items',item['id'])]:
        response=client.request('DELETE',f'/api/{kind}/{identifier}',json={'confirmation':'DELETE'})
        assert response.status_code==200,response.text
        assert response.json()['deleted']
        assert client.get(f'/api/{kind}?all_records=true').json()['total']==0
    after=client.get('/api/invoices/'+invoice['id']).json()
    assert after['customer_id'] is None and after['lines'][0]['item_id'] is None
    assert after['version']==before['version']+2
    for key in ['customer_snapshot','business_snapshot','total','taxable','cgst','sgst','paid','balance','payments','status','number']:
        assert after[key]==before[key],key
    assert after['lines']==[{**line,'item_id':None} for line in before['lines']]
    pdf=PdfReader(io.BytesIO(client.get('/api/invoices/'+invoice['id']+'/pdf').content))
    assert 'Test Beauty Store' in pdf.pages[0].extract_text()
    assert 'Original product description' in pdf.pages[0].extract_text()
    report=client.get('/api/reports/sales?start=2026-10-01&end=2026-10-31').json()
    assert report['summary']['total']==23600 and report['summary']['invoices']==1
    # A stale editor cannot put deleted master references back onto the invoice.
    assert client.put('/api/invoices/'+invoice['id'],json={**payload,'version':before['version']}).status_code==400
    saved=client.put('/api/invoices/'+invoice['id'],json={**payload,'customer_id':None,'lines':[{**payload['lines'][0],'item_id':None}],'version':after['version'],'notes':'Updated after master deletion'})
    assert saved.status_code==200,saved.text
    assert saved.json()['customer_snapshot']==before['customer_snapshot']
    assert saved.json()['total']==23600 and saved.json()['balance']==13600
    assert client.post('/api/invoices',json={**payload,'customer_id':None}).status_code==400
    with connection() as db:
        assert not db.execute('PRAGMA foreign_key_check').fetchall()
    actions=[r['action'] for r in client.get('/api/audit').json()['rows']]
    assert 'Customer deleted' in actions and 'Item deleted' in actions


@pytest.mark.parametrize('kind',['customers','items'])
def test_master_delete_requires_owner_csrf_origin_and_confirmation(client,kind):
    row=client.post('/api/'+kind,json={'name':'Disposable test record'}).json()
    url=f'/api/{kind}/{row["id"]}'
    assert client.request('DELETE',url,json={}).status_code==422
    assert client.request('DELETE',url,json={'confirmation':'yes'}).status_code==422
    csrf=client.headers.pop('X-CSRF-Token')
    assert client.request('DELETE',url,json={'confirmation':'DELETE'}).status_code==403
    client.headers['X-CSRF-Token']=csrf
    assert client.request('DELETE',url,headers={'Origin':'https://hostile.example'},json={'confirmation':'DELETE'}).status_code==403
    assert client.get('/api/'+kind).json()['total']==1
    assert client.request('DELETE',url,json={'confirmation':'DELETE'}).status_code==200
    assert client.request('DELETE',url,json={'confirmation':'DELETE'}).status_code==400
    client.post('/api/auth/logout',json={})
    assert client.request('DELETE',url,json={'confirmation':'DELETE'}).status_code==401


def test_day_month_year_in_pdf_and_excel_exports(client):
    invoice=client.post('/api/invoices',json=invoice_payload(client,invoice_date='2026-10-07',due_date='2026-11-09')).json()
    pdf=PdfReader(io.BytesIO(client.get('/api/invoices/'+invoice['id']+'/pdf').content))
    assert 'DATE: 07/10/2026' in pdf.pages[0].extract_text()
    export=load_workbook(io.BytesIO(client.get('/api/export/invoices').content)).active
    assert export['B2'].number_format=='dd/mm/yyyy'
    assert export['B2'].value.strftime('%Y-%m-%d')=='2026-10-07'
    for kind in ['sales','daily','outstanding']:
        url='/api/reports/'+kind+'?start=2026-10-01&end=2026-11-30'
        text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(client.get(url+'&format=pdf').content)).pages)
        assert '01/10/2026 to 30/11/2026' in text
        assert '07/10/2026' in text
        ws=load_workbook(io.BytesIO(client.get(url+'&format=xlsx').content)).active
        dates=[c for row in ws.iter_rows(min_row=2) for c in row if c.is_date]
        assert dates and all(c.number_format=='dd/mm/yyyy' for c in dates)
    assert client.get('/api/invoices/'+invoice['id']).json()['invoice_date']=='2026-10-07'
